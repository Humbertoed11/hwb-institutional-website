import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")
ws2 = wb["Service_Gantt_Schedule"]

print("Current Sheet 2 column widths:")
for c in range(1, 45):
    col_let = get_column_letter(c)
    w = ws2.column_dimensions[col_let].width
    if w is not None:
        print(f"  Col {col_let} ({c}): {w}")

print("\nCurrent Sheet 2 page setup:")
print("  Orientation:", ws2.page_setup.orientation)
print("  PaperSize:", ws2.page_setup.paperSize)
print("  FitToWidth:", ws2.page_setup.fitToWidth)
print("  FitToHeight:", ws2.page_setup.fitToHeight)
print("  FitToPage:", ws2.sheet_properties.pageSetUpPr.fitToPage)
