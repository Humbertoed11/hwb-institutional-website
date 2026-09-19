import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.pagebreak import Break

wb = openpyxl.load_workbook("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")

# 1. Update Commercial_Quote print_title_rows to 1:8
ws_quote = wb["Commercial_Quote"]
ws_quote.print_title_rows = '$1:$8'
print(f"Sheet Commercial_Quote print_title_rows set to: {ws_quote.print_title_rows}")

# 2. Update column widths on Commercial_Quote
col_widths_s1 = {
    1: 6.0,   # Item
    2: 26.0,  # Building Name (fits "Knights of Pythias Hall")
    3: 22.0,  # Address (fits "420 Throckmorton St")
    4: 26.0,  # Primary Floor Substrate
    5: 14.0,  # Cleanable Area (SQFT)
    6: 12.0,  # Shift Window sub-col F
    7: 12.0,  # Shift Window sub-col G
    8: 8.0,   # Crew Hrs sub-col H
    9: 8.0,   # Crew Hrs sub-col I
    10: 14.0, # Service sub-col J
    11: 14.0  # Service sub-col K
}
for col_idx, width in col_widths_s1.items():
    ws_quote.column_dimensions[get_column_letter(col_idx)].width = width

# Set Day columns L to AT (Cols 12 to 46) to 4.5 width (comfortably uncrowded!)
for day_col in range(12, 47):
    ws_quote.column_dimensions[get_column_letter(day_col)].width = 4.5

# 3. Update Service_Gantt_Schedule (Sheet 2)
ws_gantt = wb["Service_Gantt_Schedule"]
ws_gantt.print_title_rows = '$1:$4' # Masthead, Subtitle, Meta, Legend
ws_gantt.column_dimensions['A'].width = 6.0   # Item
ws_gantt.column_dimensions['B'].width = 26.0  # Building Name
ws_gantt.column_dimensions['C'].width = 14.0  # Total SQF
ws_gantt.column_dimensions['D'].width = 22.0  # Shift Window
ws_gantt.column_dimensions['E'].width = 15.0  # Crew Hrs
for col_idx in range(6, 41):
    ws_gantt.column_dimensions[get_column_letter(col_idx)].width = 4.5

# Ensure print margins and fit properties
for ws in [ws_quote, ws_gantt]:
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35

ws_quote.page_setup.fitToHeight = 0
ws_gantt.page_setup.fitToHeight = 1

wb.save("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")
print("SUCCESS: Updated SUNDANCE-QUOTE-gantt.xlsx with fitted uncrowded columns and $1:$8 repeating header!")
