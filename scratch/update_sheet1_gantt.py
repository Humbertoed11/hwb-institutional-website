import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")
ws1 = wb["Commercial_Quote"]

# Styling definitions
font_sec = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
font_bold = Font(name="Calibri", size=9.5, bold=True, color="0F172A")
font_reg = Font(name="Calibri", size=9.5, color="334155")
font_hdr = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
font_badge = Font(name="Calibri", size=7, bold=True, color="FFFFFF")

fill_navy = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
fill_subtle = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_weekend = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

fill_deep = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
fill_seal = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
fill_qtr = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
fill_stair = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")

align_c = Alignment(horizontal="center", vertical="center")
align_l = Alignment(horizontal="left", vertical="center")
align_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_side = Side(style="thin", color="CBD5E1")
border_box = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
border_tot = Border(top=Side(style="thin", color="0F172A"), bottom=Side(style="double", color="0F172A"), left=thin_side, right=thin_side)

# Update Row 79: Legend Bar on Sheet 1
ws1.row_dimensions[79].height = 20
ws1.merge_cells("A79:D79")
ws1["A79"] = "LEGEND & COLOR CODES:"
ws1["A79"].font = font_bold
ws1["A79"].fill = fill_subtle
ws1["A79"].alignment = align_c

legend_sec5 = [
    ("E79:F79", "DP: Phase 1 Deep Clean", fill_deep),
    ("G79:H79", "SL: Modular Sealer", fill_seal),
    ("I79:J79", "QT: Quarterly Care", fill_qtr),
    ("K79", "ST: Stairs", fill_stair)
]
for crange, ltext, lfill in legend_sec5:
    if ":" in crange:
        ws1.merge_cells(crange)
        c_top = ws1[crange.split(":")[0]]
    else:
        c_top = ws1[crange]
    c_top.value = ltext
    c_top.font = font_badge
    c_top.fill = lfill
    c_top.alignment = align_c

# Also add Legend above timeline columns (L79 to AT79)
ws1.merge_cells("L79:R79")
ws1["L79"] = "DP: Phase 1 Deep Clean"
ws1["L79"].font = font_badge
ws1["L79"].fill = fill_deep
ws1["L79"].alignment = align_c

ws1.merge_cells("S79:Y79")
ws1["S79"] = "SL: Modular Sealer (BulletProof®)"
ws1["S79"].font = font_badge
ws1["S79"].fill = fill_seal
ws1["S79"].alignment = align_c

ws1.merge_cells("Z79:AF79")
ws1["Z79"] = "QT: Track A Quarterly Care"
ws1["Z79"].font = font_badge
ws1["Z79"].fill = fill_qtr
ws1["Z79"].alignment = align_c

ws1.merge_cells("AG79:AM79")
ws1["AG79"] = "ST: Stair Detailing (500 Main)"
ws1["AG79"].font = font_badge
ws1["AG79"].fill = fill_stair
ws1["AG79"].alignment = align_c

ws1.merge_cells("AN79:AT79")
ws1["AN79"] = "⏰ 100% Flexible Shift Timing"
ws1["AN79"].font = Font(name="Calibri", size=7.5, bold=True, color="16A34A")
ws1["AN79"].fill = fill_subtle
ws1["AN79"].alignment = align_c

for c in range(1, 47):
    ws1.cell(row=79, column=c).border = border_box

# Update Data Rows 83 to 93 with 2-letter badges (DP, SL, QT, ST)
schedule_data = [
    (83, [13], [17], [27], []),
    (84, [14], [17], [28], []),
    (85, [15], [17], [29], []),
    (86, [16], [17], [30], []),
    (87, [19], [24], [33], [34]),
    (88, [20], [24], [34], []),
    (89, [21], [24], [35], []),
    (90, [22], [24], [36], []),
    (91, [23], [24], [37], []),
    (92, [23], [24], [37], []),
    (93, [26], [27], [40], [])
]

for row_idx, p1_cols, seal_cols, rec_cols, stair_cols in schedule_data:
    for day_i in range(1, 36):
        col = day_i + 11
        cell = ws1.cell(row=row_idx, column=col)
        cell.border = border_box
        
        if col in p1_cols:
            cell.fill = fill_deep
            cell.value = "DP"
            cell.font = font_badge
            cell.alignment = align_c
        elif col in seal_cols:
            cell.fill = fill_seal
            cell.value = "SL"
            cell.font = font_badge
            cell.alignment = align_c
        elif col in rec_cols:
            cell.fill = fill_qtr
            cell.value = "QT"
            cell.font = font_badge
            cell.alignment = align_c
        elif col in stair_cols:
            cell.fill = fill_stair
            cell.value = "ST"
            cell.font = font_badge
            cell.alignment = align_c
        elif (col - 12) % 7 in [5, 6]:
            cell.fill = fill_weekend
            cell.value = None
        else:
            cell.fill = PatternFill(fill_type=None)
            cell.value = None

# Update column widths for Columns L to AT (Cols 12 to 46) to 2.8!
for col_idx in range(12, 47):
    ws1.column_dimensions[get_column_letter(col_idx)].width = 2.8

wb.save("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")
print("SUCCESS: Section 5.0 on Sheet 1 (Commercial_Quote) updated to Option 1 (DP, SL, QT, ST, width 2.8)!")
