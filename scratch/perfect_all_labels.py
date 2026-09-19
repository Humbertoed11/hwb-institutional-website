import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, range_boundaries

wb = openpyxl.load_workbook("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")

# 1. Commercial_Quote
ws1 = wb["Commercial_Quote"]

# Width adjustments on Sheet 1:
ws1.column_dimensions['A'].width = 5.5   # Item
ws1.column_dimensions['B'].width = 25.0  # Building Name
ws1.column_dimensions['C'].width = 22.0  # Address
ws1.column_dimensions['D'].width = 24.0  # Primary Floor Substrate
ws1.column_dimensions['E'].width = 14.0  # Cleanable Area (SQFT)
ws1.column_dimensions['F'].width = 13.0  # Phase 1 Rate
ws1.column_dimensions['G'].width = 18.0  # Phase 1 deep clean (expanded to 18 for zero wrapping inside lines)
ws1.column_dimensions['H'].width = 20.0  # Option B Stair Care / Quarterly Rate (expanded to 20 to fit "$3.00/Stp + $21/Lnd" and "N/A (Excluded)")
ws1.column_dimensions['I'].width = 16.0  # Track A: Quarterly
ws1.column_dimensions['J'].width = 14.0  # Semi-Annual Rate
ws1.column_dimensions['K'].width = 20.0  # Track B: Semi-Annual / Available Per Hour (expanded to 20 for zero clipping)

# Header Row 17 (Table 1) height:
ws1.row_dimensions[17].height = 46

# Header Row 32 (Table 2) height:
ws1.row_dimensions[32].height = 46

# Table 3 Term Rows (Rows 48 to 54) height:
for r in range(48, 55):
    ws1.row_dimensions[r].height = 50

# 2. Service_Gantt_Schedule
ws2 = wb["Service_Gantt_Schedule"]

for rng in list(ws2.merged_cells.ranges):
    if rng.min_row <= 3 <= rng.max_row:
        ws2.unmerge_cells(str(rng))

font_bold = Font(name="Calibri", size=9, bold=True, color="0F172A")
font_reg = Font(name="Calibri", size=9, color="334155")
thin_side = Side(style="thin", color="CBD5E1")
box_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)

ws2.row_dimensions[3].height = 18

ws2.merge_cells("A3:C3")
ws2["A3"] = "Base Dispatch:"
ws2["A3"].font = font_bold
ws2["A3"].alignment = Alignment(horizontal="center", vertical="center")

ws2.merge_cells("D3:G3")
ws2["D3"] = "Arlington, TX (Post-7:00 PM)"
ws2["D3"].font = font_reg
ws2["D3"].alignment = Alignment(horizontal="left", vertical="center")

ws2.merge_cells("H3:J3")
ws2["H3"] = "Start Date:"
ws2["H3"].font = font_bold
ws2["H3"].alignment = Alignment(horizontal="center", vertical="center")

ws2.merge_cells("K3:N3")
ws2["K3"] = "Sep 01, 2026"
ws2["K3"].font = font_bold
ws2["K3"].alignment = Alignment(horizontal="center", vertical="center")

ws2.merge_cells("O3:R3")
ws2["O3"] = "Assigned Crew:"
ws2["O3"].font = font_bold
ws2["O3"].alignment = Alignment(horizontal="center", vertical="center")

ws2.merge_cells("S3:X3")
ws2["S3"] = "2 Specialized Technicians"
ws2["S3"].font = font_reg
ws2["S3"].alignment = Alignment(horizontal="left", vertical="center")

ws2.merge_cells("Y3:AC3")
ws2["Y3"] = "Mobilization:"
ws2["Y3"].font = font_bold
ws2["Y3"].alignment = Alignment(horizontal="center", vertical="center")

ws2.merge_cells("AD3:AH3")
ws2["AD3"] = "10 Dedicated Shifts"
ws2["AD3"].font = font_reg
ws2["AD3"].alignment = Alignment(horizontal="left", vertical="center")

ws2.merge_cells("AI3:AN3")
ws2["AI3"] = "⏰ 100% Flexible Shift Timing"
ws2["AI3"].font = Font(name="Calibri", size=8.5, bold=True, color="16A34A")
ws2["AI3"].alignment = Alignment(horizontal="center", vertical="center")

for c in range(1, 41):
    ws2.cell(row=3, column=c).border = box_border

# 3. Cost_Underwriting_Audit
ws3 = wb["Cost_Underwriting_Audit"]
for rng in list(ws3.merged_cells.ranges):
    if rng.min_row <= 16 <= rng.max_row:
        ws3.unmerge_cells(str(rng))

ws3.merge_cells("A16:D16")
ws3["A16"] = "TOTAL PER QUARTERLY PASS (10 Dedicated Night Shifts)"
ws3["A16"].font = Font(name="Calibri", size=9.5, bold=True)
ws3["A16"].alignment = Alignment(horizontal="center", vertical="center")

wb.save("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")
print("SUCCESS: Perfected all labels across all sheets!")
