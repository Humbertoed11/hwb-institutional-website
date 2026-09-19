import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook("HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx")
ws1 = wb["Commercial_Quote"]

print("Current Sheet 1 max column:", ws1.max_column) # 46 (AT)

# If we extend:
# 1. Title Block: A1:AT1
# 2. Subtitle Block: A2:AT2
# 3. Meta Info: A:AT
# 4. Data Integrity: A8:AT8
# 5. Executive Summary: A10:AT10, A11:AT11, etc.
# 6. Table 1, Table 2, Table 3, Table 4:
# Let's inspect how Table 1 columns can map across A to AT!

print("Done inspecting.")
