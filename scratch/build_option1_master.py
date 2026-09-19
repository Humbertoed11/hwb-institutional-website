import os
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.defined_name import DefinedName

def generate_option1_workbook():
    target_path = "HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx"
    
    img_paths = {
        "plaza": "scratch/thumb_plaza.jpg",
        "chase": "scratch/thumb_chase.jpg",
        "petroleum": "scratch/thumb_petroleum.jpg",
        "schwarz": "scratch/thumb_schwarz.jpg"
    }

    wb = openpyxl.Workbook()

    # Define Named Ranges
    defined_name_start = DefinedName('Project_Start', attr_text='Service_Gantt_Schedule!$G$4')
    defined_name_week = DefinedName('Display_Week', attr_text='Service_Gantt_Schedule!$G$5')
    wb.defined_names.add(defined_name_start)
    wb.defined_names.add(defined_name_week)

    # Styling Definitions
    font_title = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=10.5, bold=True, color="1E3A8A")
    font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=9, bold=True, color="0F172A")
    font_bold_sm = Font(name="Calibri", size=8.5, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=9, color="334155")
    font_regular_sm = Font(name="Calibri", size=8, color="334155")
    font_alert = Font(name="Calibri", size=8.5, bold=True, color="991B1B")
    font_badge = Font(name="Calibri", size=7.5, bold=True, color="FFFFFF")

    fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header_opt = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_subtle = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_alert = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    fill_green = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    fill_weekend = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    fill_p1_deep_clean = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    fill_p2_sealer = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
    fill_p3_quarterly = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    fill_opt_stairs = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_header_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_desc_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_side = Side(style="thin", color="CBD5E1")
    box_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    total_border = Border(top=Side(style="thin", color="0F172A"), bottom=Side(style="double", color="0F172A"), left=thin_side, right=thin_side)

    # =============================================================
    # SHEET 1: Commercial_Quote (Strictly 11 Columns A to K: 100% Full Scale)
    # =============================================================
    ws_quote = wb.active
    ws_quote.title = "Commercial_Quote"
    ws_quote.views.sheetView[0].showGridLines = True

    # Title Block (Rows 1 & 2)
    ws_quote.merge_cells("A1:K1")
    ws_quote["A1"] = "HWB CLEANING SERVICES LLC"
    ws_quote["A1"].font = font_title
    ws_quote["A1"].fill = fill_title
    ws_quote["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_quote.row_dimensions[1].height = 24

    ws_quote.merge_cells("A2:K2")
    ws_quote["A2"] = "COMMERCIAL PROPOSAL: BOSANNA LLC — 11-BUILDING COMMERCIAL PORTFOLIO"
    ws_quote["A2"].font = Font(name="Calibri", size=9.5, bold=True, color="93C5FD")
    ws_quote["A2"].fill = fill_title
    ws_quote["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_quote.row_dimensions[2].height = 16

    ws_quote.row_dimensions[3].height = 4

    # Meta Info Box (Rows 4 to 7) - Columns A to K
    meta_info = [
        ("Client / Prime:", "Bosanna LLC (Attn: Angelica Hudgins)", "Proposal ID:", "HWB-BID-2026-FW01 (v3.8.0)"),
        ("End Client Portfolio:", "Bosanna LLC (11 Landmark Buildings)", "Effective Date:", "September 04, 2026"),
        ("Central Office:", "425 Houston Street, Suite 250, Fort Worth, TX 76102", "Dispatch Origin:", "Arlington, TX"),
        ("Primary Scope:", "11 Ground-Floor Commercial Lobbies (11,153 Cleanable SF)", "Subcontractor Role:", "Specialized Floor Care Crew (2 Techs)")
    ]
    for idx, (lbl1, val1, lbl2, val2) in enumerate(meta_info, start=4):
        ws_quote.row_dimensions[idx].height = 18
        ws_quote.merge_cells("A{0}:B{0}".format(idx))
        ws_quote["A{0}".format(idx)] = lbl1
        ws_quote["A{0}".format(idx)].font = font_bold
        ws_quote["A{0}".format(idx)].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws_quote.merge_cells("C{0}:F{0}".format(idx))
        ws_quote["C{0}".format(idx)] = val1
        ws_quote["C{0}".format(idx)].font = font_regular
        ws_quote["C{0}".format(idx)].alignment = Alignment(horizontal="left", vertical="center")

        ws_quote.merge_cells("G{0}:H{0}".format(idx))
        ws_quote["G{0}".format(idx)] = lbl2
        ws_quote["G{0}".format(idx)].font = font_bold
        ws_quote["G{0}".format(idx)].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws_quote.merge_cells("I{0}:K{0}".format(idx))
        ws_quote["I{0}".format(idx)] = val2
        ws_quote["I{0}".format(idx)].font = font_regular
        ws_quote["I{0}".format(idx)].alignment = Alignment(horizontal="left", vertical="center")

        for c in range(1, 12):
            cell = ws_quote.cell(row=idx, column=c)
            cell.border = box_border
            if c in [1, 2, 7, 8]:
                cell.fill = fill_zebra

    # Row 8: Data Integrity Alert Notice
    ws_quote.merge_cells("A8:K8")
    ws_quote["A8"] = "DATA INTEGRITY NOTICE (ISO 9001 Clause 8.2.2): Cleanable square footages are reconciled to on-site empirical field measurements (11,153 Cleanable SF)."
    ws_quote["A8"].font = font_alert
    ws_quote["A8"].fill = fill_alert
    ws_quote["A8"].alignment = align_center
    ws_quote.row_dimensions[8].height = 18
    for c in range(1, 12):
        ws_quote.cell(row=8, column=c).border = box_border

    ws_quote.row_dimensions[9].height = 6

    # Executive Summary Block (Rows 10 to 14)
    ws_quote.merge_cells("A10:K10")
    ws_quote["A10"] = "EXECUTIVE SUMMARY: CLIENT NEEDS, GOALS & OUR PROPOSED PLAN"
    ws_quote["A10"].font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    ws_quote["A10"].fill = fill_header
    ws_quote["A10"].alignment = align_center
    ws_quote.row_dimensions[10].height = 20

    ws_quote.merge_cells("A11:K11")
    ws_quote["A11"] = "1. Current Situation & Needs: Bosanna LLC oversees facilities operations for its commercial portfolio across 11 landmark downtown buildings in Sundance Square. Heavy daily foot traffic, grit, and rain have dulled entrance lobbies, settling dirt into stone and grout pores. As specialized floor care subcontractor to Bosanna LLC, HWB ensures all lobbies project an immaculate first impression."
    ws_quote["A11"].font = font_regular
    ws_quote["A11"].alignment = align_desc_wrap
    ws_quote.row_dimensions[11].height = 28

    ws_quote.merge_cells("A12:K12")
    ws_quote["A12"] = "2. Client Goals: (1) Strip away old ground-in soil to reveal true floor condition and existing damage; (2) Seal stone and grout pores to prevent coffee, drink, and grease stains; (3) Set up a predictable maintenance routine (quarterly/semi-annual) within budget; (4) Ensure zero daytime tenant disruption."
    ws_quote["A12"].font = font_regular
    ws_quote["A12"].alignment = align_desc_wrap
    ws_quote.row_dimensions[12].height = 28

    ws_quote.merge_cells("A13:K13")
    ws_quote["A13"] = "3. Proposed 3-Phase Solution: (1) Phase 1 Deep Clean Reset (rotary machine scrub & edge extraction); (2) Phase 2 Sub-Surface Sealer (STONETECH® BulletProof® oil/water barrier); (3) Phase 3 Scheduled Care via Track A (Quarterly) or Track B (Semi-Annual) recurring programs."
    ws_quote["A13"].font = font_regular
    ws_quote["A13"].alignment = align_desc_wrap
    ws_quote.row_dimensions[13].height = 26

    ws_quote.merge_cells("A14:K14")
    ws_quote["A14"] = "💰 Bosanna LLC Subcontractor Direct Partnership Advantage: As a specialized floor care subcontractor to Bosanna LLC, HWB has applied a full 25% direct discount on square footage rates across all 11 properties—providing a direct 20¢ per sq. ft. discount on Phase 1 deep cleaning (lowering the rate from $0.80 down to $0.60/SF) and a total 27.75¢ per sq. ft. discount across complete initial reset & sealing (reducing from $1.11 down to $0.8325/SF; Total Initial Reset: $9,685.30 complete vs. $12,913.74 regular, saving Bosanna $3,228.44). Operations run on flexible night shifts post-7:00 PM with zero tenant disruption."
    ws_quote["A14"].font = Font(name="Calibri", size=8.5, bold=True, color="14532D")
    ws_quote["A14"].fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    ws_quote["A14"].alignment = align_desc_wrap
    ws_quote.row_dimensions[14].height = 42

    for r_idx in range(10, 15):
        for c in range(1, 12):
            ws_quote.cell(row=r_idx, column=c).border = box_border

    ws_quote.row_dimensions[15].height = 6

    # -------------------------------------------------------------
    # SECTION 1.0: Commercial Pricing Schedule (Table 1: Rows 16 to 29)
    # -------------------------------------------------------------
    ws_quote.merge_cells("A16:K16")
    ws_quote["A16"] = "1.0 COMMERCIAL PRICING SCHEDULE: INITIAL RESET & RECURRING MAINTENANCE PROGRAMS"
    ws_quote["A16"].font = font_section
    ws_quote.row_dimensions[16].height = 20

    t1_headers = [
        "Item",
        "Building Name",
        "Address",
        "Primary floor system",
        "Cleanable Area\n(SQFT)",
        "Phase 1 Rate\n(Per SF)",
        "Phase 1 deep clean\n(expose damage to\nconsider options)",
        "Quarterly Rate\n(Per SF)",
        "Track A: Quarterly\n(Every 3 Months)",
        "Semi-Annual Rate\n(Per SF)",
        "Track B: Semi-Annual\n(Every 6 Months)"
    ]
    ws_quote.row_dimensions[17].height = 32
    for col_idx, htext in enumerate(t1_headers, start=1):
        cell = ws_quote.cell(row=17, column=col_idx, value=htext)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_header_wrap
        cell.border = box_border

    portfolio_data = [
        (1, "Chase Bank Building", "420 Throckmorton St", "Polished Terrazzo", 2126, 0.60, 0.405, 0.51),
        (2, "The Westbrook", "425 Houston St", "Polished Terrazzo", 1520, 0.60, 0.405, 0.51),
        (3, "The Carnegie", "421 W 3rd St", "Polished Terrazzo", 1175, 0.60, 0.405, 0.51),
        (4, "The Cassidy", "407 Throckmorton St", "Polished Terrazzo", 1134, 0.60, 0.405, 0.51),
        (5, "Burk Burnett Building", "500 Main St", "Ceramic Mosaic Tile", 1108, 0.60, 0.405, 0.51),
        (6, "Sanger Lofts", "222 W 4th St", "Polished Terrazzo", 1071, 0.60, 0.405, 0.51),
        (7, "Petroleum Building", "210 W 6th St", "Polished Marble Tile", 860, 0.60, 0.405, 0.51),
        (8, "The Commerce Building", "420 Commerce St", "Terrazzo Tile", 840, 0.60, 0.405, 0.51),
        (9, "Virtuoso Building", "505 Main St", "Polished Marble Tile", 560, 0.60, 0.405, 0.51),
        (10, "Knights of Pythias Hall", "109-110 E 3rd St", "Ceramic Mosaic Tile", 414, 0.60, 0.405, 0.51),
        (11, "Plaza Hotel Building", "303 Main St", "Commercial Quarry Tile", 345, 0.60, 0.405, 0.51)
    ]

    for idx, item in enumerate(portfolio_data, start=18):
        r_num, b_name, b_addr, b_sub, b_sf, r_mo, r_qtr, r_semi = item
        ws_quote.row_dimensions[idx].height = 16
        fill_row = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)

        ws_quote.cell(row=idx, column=1, value=r_num).alignment = align_center
        ws_quote.cell(row=idx, column=2, value=b_name).alignment = align_left
        ws_quote.cell(row=idx, column=3, value=b_addr).alignment = align_left
        ws_quote.cell(row=idx, column=4, value=b_sub).alignment = align_left

        c_sf = ws_quote.cell(row=idx, column=5, value=b_sf)
        c_sf.number_format = "#,##0"
        c_sf.alignment = align_center

        c_rmo = ws_quote.cell(row=idx, column=6, value=r_mo)
        c_rmo.number_format = "$#,##0.00"
        c_rmo.alignment = align_right

        c_mo = ws_quote.cell(row=idx, column=7, value="=MAX(E{0},600)*F{0}".format(idx))
        c_mo.number_format = "$#,##0.00"
        c_mo.alignment = align_right

        c_rq = ws_quote.cell(row=idx, column=8, value=r_qtr)
        c_rq.number_format = "$#,##0.00"
        c_rq.alignment = align_right

        c_qtr = ws_quote.cell(row=idx, column=9, value="=MAX(E{0},600)*H{0}".format(idx))
        c_qtr.number_format = "$#,##0.00"
        c_qtr.alignment = align_right

        c_rs = ws_quote.cell(row=idx, column=10, value=r_semi)
        c_rs.number_format = "$#,##0.00"
        c_rs.alignment = align_right

        c_semi = ws_quote.cell(row=idx, column=11, value="=MAX(E{0},600)*J{0}".format(idx))
        c_semi.number_format = "$#,##0.00"
        c_semi.alignment = align_right

        for c in range(1, 12):
            cell = ws_quote.cell(row=idx, column=c)
            cell.font = font_regular
            cell.border = box_border
            if fill_row.fill_type:
                cell.fill = fill_row

    # Total Row Table 1 (Row 29)
    tot_r1 = 29
    ws_quote.row_dimensions[tot_r1].height = 20
    ws_quote.merge_cells("A{0}:D{0}".format(tot_r1))
    ws_quote["A{0}".format(tot_r1)] = "GRAND TOTAL SQUARE FOOTAGE (11 Lobbies)"
    ws_quote["A{0}".format(tot_r1)].font = font_bold
    ws_quote["A{0}".format(tot_r1)].alignment = align_center

    c_tot_sf1 = ws_quote.cell(row=tot_r1, column=5, value="=SUM(E18:E28)")
    c_tot_sf1.number_format = '#,##0 " SQF"'
    c_tot_sf1.alignment = align_center

    ws_quote.cell(row=tot_r1, column=6, value="$0.60 avg").alignment = align_center
    c_tot_mo = ws_quote.cell(row=tot_r1, column=7, value="=SUM(G18:G28)")
    c_tot_mo.number_format = "$#,##0.00"
    c_tot_mo.alignment = align_right

    ws_quote.cell(row=tot_r1, column=8, value="$0.405 avg").alignment = align_center
    c_tot_qtr = ws_quote.cell(row=tot_r1, column=9, value="=SUM(I18:I28)")
    c_tot_qtr.number_format = "$#,##0.00"
    c_tot_qtr.alignment = align_right

    ws_quote.cell(row=tot_r1, column=10, value="$0.51 avg").alignment = align_center
    c_tot_semi = ws_quote.cell(row=tot_r1, column=11, value="=SUM(K18:K28)")
    c_tot_semi.number_format = "$#,##0.00"
    c_tot_semi.alignment = align_right

    for c in range(1, 12):
        cell = ws_quote.cell(row=tot_r1, column=c)
        cell.font = font_bold
        cell.fill = fill_subtle
        cell.border = total_border

    # Page Break 1 (After Row 29 -> Ends Page 1)
    ws_quote.row_breaks.append(Break(id=29))

    # -------------------------------------------------------------
    # PAGE 2: SECTION 2.0 (Table 2) & SECTION 3.0 (Table 3)
    # -------------------------------------------------------------
    ws_quote.row_dimensions[30].height = 6
    ws_quote.merge_cells("A31:K31")
    ws_quote["A31"] = "2.0 ADDITIONAL CLIENT OPTIONS: PENETRATING SEALER, STAIRCARE, DIAMOND RESTORATION & ART HANDLING"
    ws_quote["A31"].font = font_section
    ws_quote.row_dimensions[31].height = 20

    t2_headers = [
        ("A32", "Item"),
        ("B32", "Building Name"),
        ("C32", "Address"),
        ("D32", "Primary floor system"),
        ("E32", "Cleanable Area\n(SQFT)"),
        ("F32:G32", "Option A: Sealer\n(BulletProof®)\n$0.2325 / SF"),
        ("H32", "Option B: Stair Care\n(Unit Rate Add-On\n500 Main St Only)*"),
        ("I32:J32", "Option C: Capital Diamond Restoration\n(Restores Mirror Shine • Marble & Terrazzo Only)\n$1.9875 / SF"),
        ("K32", "Option D: Furniture Handling\n(Hourly T&M)\nAvailable Per Hour")
    ]
    ws_quote.row_dimensions[32].height = 32
    for crange, htext in t2_headers:
        if ":" in crange:
            ws_quote.merge_cells(crange)
            c_top = ws_quote[crange.split(":")[0]]
            c_start, _, c_end, _ = range_boundaries(crange)
            for c in range(c_start, c_end + 1):
                ws_quote.cell(row=32, column=c).border = box_border
                ws_quote.cell(row=32, column=c).fill = fill_header_opt
        else:
            c_top = ws_quote[crange]
            c_top.border = box_border
            c_top.fill = fill_header_opt
        c_top.value = htext
        c_top.font = font_header
        c_top.alignment = align_header_wrap

    for idx, item in enumerate(portfolio_data, start=33):
        num, name, addr, sub, sf, _, _, _ = item
        base_row = idx - 15
        ws_quote.row_dimensions[idx].height = 16
        fill_row = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)

        ws_quote.cell(row=idx, column=1, value=num).alignment = align_center
        ws_quote.cell(row=idx, column=2, value=name).alignment = align_left
        ws_quote.cell(row=idx, column=3, value=addr).alignment = align_left
        ws_quote.cell(row=idx, column=4, value=sub).alignment = align_left

        c_sf2 = ws_quote.cell(row=idx, column=5, value="=E{0}".format(base_row))
        c_sf2.number_format = "#,##0"
        c_sf2.alignment = align_center

        ws_quote.merge_cells("F{0}:G{0}".format(idx))
        c_seal = ws_quote.cell(row=idx, column=6, value="=MAX(E{0},600)*0.2325".format(idx))
        c_seal.number_format = "$#,##0.00"
        c_seal.alignment = align_right

        c_stv = ws_quote.cell(row=idx, column=8)
        if num == 5:
            c_stv.value = "$3/Stp + $21/Lnd"
            c_stv.font = Font(name="Calibri", size=8.5, bold=True, color="16A34A")
        else:
            c_stv.value = "N/A (Excluded)"
            c_stv.font = Font(name="Calibri", size=8.5, italic=True, color="64748B")
        c_stv.alignment = align_center

        ws_quote.merge_cells("I{0}:J{0}".format(idx))
        is_grindable = sub in ["Polished Terrazzo", "Terrazzo Tile", "Polished Marble Tile"]
        c_rst = ws_quote.cell(row=idx, column=9)
        if is_grindable:
            c_rst.value = "=MAX(E{0},600)*1.9875".format(idx)
            c_rst.number_format = "$#,##0.00"
            c_rst.alignment = align_right
        else:
            c_rst.value = "N/A (Exempt)"
            c_rst.font = Font(name="Calibri", size=8.5, italic=True, color="64748B")
            c_rst.alignment = align_center

        c_fur = ws_quote.cell(row=idx, column=11, value="Available Per Hour")
        c_fur.alignment = align_center

        for c in range(1, 12):
            cell = ws_quote.cell(row=idx, column=c)
            if c not in [8, 9, 10]:
                cell.font = font_regular
            cell.border = box_border
            if fill_row.fill_type:
                cell.fill = fill_row

    # Total Row Table 2 (Row 44)
    tot_r2 = 44
    ws_quote.row_dimensions[tot_r2].height = 20
    ws_quote.merge_cells("A{0}:D{0}".format(tot_r2))
    ws_quote["A{0}".format(tot_r2)] = "GRAND TOTAL SQUARE FOOTAGE (Additional Options)"
    ws_quote["A{0}".format(tot_r2)].font = font_bold
    ws_quote["A{0}".format(tot_r2)].alignment = align_center

    c_tot_sf2 = ws_quote.cell(row=tot_r2, column=5, value="=SUM(E33:E43)")
    c_tot_sf2.number_format = '#,##0 " SQF"'
    c_tot_sf2.alignment = align_center

    ws_quote.merge_cells("F{0}:G{0}".format(tot_r2))
    c_tot_sl2 = ws_quote.cell(row=tot_r2, column=6, value="=SUM(F33:F43)")
    c_tot_sl2.number_format = "$#,##0.00"
    c_tot_sl2.alignment = align_right

    ws_quote.cell(row=tot_r2, column=8, value="Unit Rate Only").alignment = align_center

    ws_quote.merge_cells("I{0}:J{0}".format(tot_r2))
    c_tot_rst2 = ws_quote.cell(row=tot_r2, column=9, value="=SUM(I33:I43)")
    c_tot_rst2.number_format = "$#,##0.00"
    c_tot_rst2.alignment = align_right

    ws_quote.cell(row=tot_r2, column=11, value="Available Per Hour").alignment = align_center

    for c in range(1, 12):
        cell = ws_quote.cell(row=tot_r2, column=c)
        cell.font = font_bold
        cell.fill = fill_subtle
        cell.border = total_border

    ws_quote.row_dimensions[45].height = 6

    # SECTION 3.0: Program Terms & Scope (Table 3: Rows 46 to 54)
    ws_quote.merge_cells("A46:K46")
    ws_quote["A46"] = "3.0 COMMERCIAL PROGRAM EXECUTIVE SUMMARY & BILLING TERMS"
    ws_quote["A46"].font = font_section
    ws_quote.row_dimensions[46].height = 20

    t3_headers = [
        ("A47:B47", "Service Phase / Scope Module"),
        ("C47:F47", "Scope of work"),
        ("G47", "Rate Basis"),
        ("H47:I47", "Portfolio Invoiced Investment"),
        ("J47:K47", "Billing Frequency & Program Terms")
    ]
    ws_quote.row_dimensions[47].height = 18
    for crange, htext in t3_headers:
        ws_quote.merge_cells(crange)
        c_top = ws_quote[crange.split(":")[0]]
        c_top.value = htext
        c_top.font = font_header
        c_top.fill = fill_header_opt
        c_top.alignment = align_center
        c_start, _, c_end, _ = range_boundaries(crange)
        for c in range(c_start, c_end + 1):
            ws_quote.cell(row=47, column=c).border = box_border
            ws_quote.cell(row=47, column=c).fill = fill_header_opt

    summary_rows = [
        ("Phase 1 deep clean (expose damage to consider options)", "Rotary machine scrub, deep cleaning solution, edge hand scraping, and extraction across 11 lobbies (600 SF min baseline). Strips surface dirt down to bare floor to expose existing damage for mutual review; does not grind stone or remove scratches. Reflects 25% direct portfolio discount (20¢/SF discount off $0.80 regular rate).", "$0.60 / SF", "$6,980.40 complete", "One-Time Baseline Service"),
        ("Phase 2: Modular Sub-Surface Sealer (STONETECH® BulletProof®)", "Deep penetrating sealer applied after cleaning into stone and grout lines (600 SF baseline). Repels water and oil spills; does not create a wax film or shine, and does not stop acid etch marks on polished marble. Reflects 25% direct portfolio discount (7.75¢/SF discount off $0.31 regular rate).", "$0.2325 / SF", "$2,704.90 complete", "Post-Clean Protection\n(12–24 Month Lifecycle)"),
        ("Phase 3 (Track A): Quarterly Maintenance (Recommended Program)", "Recurring 90-day maintenance scrub across 8 Terrazzo/Marble lobbies; deep brush scrub and grout extraction across 3 Mosaic/Quarry lobbies (600 SF min). Routine cleaning only; does not include diamond grinding. Reflects 25% direct portfolio discount (13.5¢/SF discount off $0.54 regular rate).", "$0.405 / SF", "$4,711.77 / pass", "Every 3 Months\n(4x / Year • Recommended)"),
        ("Phase 3 (Track B): Every Six Months Deep Restorative Scrub", "Intensive 180-day deep scrub, heavy cleaner dwell time, and machine extraction across 11 lobbies (600 SF min baseline). Removes 6 months of ground-in traffic dirt; does not include diamond grinding. Reflects 25% direct portfolio discount (17¢/SF discount off $0.68 regular rate).", "$0.51 / SF", "$5,933.34 / pass", "Every 6 Months\n(2x / Year Semi-Annual)"),
        ("Add-On Option B: Dedicated Marble Staircase Care (500 Main St Only)", "Step machine scrub ($3.00/step), riser wipe-down, brass nosing polish, landing deep clean ($21.00/landing), and sealer at 500 Main St only. Billed strictly by actual step count on site; excludes repairs and other 10 buildings.", "$3/Stp + $21/Lnd", "Unit Rate Only\n(T&M Count)", "Optional Per Service Pass\n(500 Main St Only)"),
        ("Add-On Option C: Capital Diamond Restoration", "Multi-step diamond grinding and polishing with wall protection across 8 stone lobbies (9,326 Billing SF). The only service that cuts out deep scratches and restores a high-gloss shine; ceramic and quarry tile are exempt. Reflects 25% direct portfolio discount (66.25¢/SF discount off $2.65 regular rate).", "$1.9875 / SF", "$18,535.43 complete", "As-Needed Capital Project\n(8 Stone Lobbies)"),
        ("Add-On Option D: Furniture & Artwork Handling", "Building staff clears and replaces lobby furniture before and after service. Crew moving assistance is available upon written request on an hourly basis; HWB is not responsible for unmoved furniture.", "Available Per Hour", "Hourly T&M\n(Upon Request)", "Client-Managed Default\n(Hourly Moving Option)")
    ]

    for s_idx, s_item in enumerate(summary_rows, start=48):
        ws_quote.row_dimensions[s_idx].height = 50
        p_name, p_desc, p_rate, p_pass, p_bill = s_item
        fill_r = fill_green if s_idx == 48 else (fill_zebra if s_idx % 2 == 0 else PatternFill(fill_type=None))

        ws_quote.merge_cells("A{0}:B{0}".format(s_idx))
        c_p = ws_quote["A{0}".format(s_idx)]
        c_p.value = p_name
        c_p.font = font_bold_sm
        c_p.alignment = align_desc_wrap

        ws_quote.merge_cells("C{0}:F{0}".format(s_idx))
        c_d = ws_quote["C{0}".format(s_idx)]
        c_d.value = p_desc
        c_d.font = font_regular_sm
        c_d.alignment = align_desc_wrap

        c_r = ws_quote.cell(row=s_idx, column=7, value=p_rate)
        c_r.font = font_bold_sm
        c_r.alignment = align_center

        ws_quote.merge_cells("H{0}:I{0}".format(s_idx))
        c_pass = ws_quote.cell(row=s_idx, column=8, value=p_pass)
        c_pass.font = font_bold_sm if ("complete" in p_pass or "pass" in p_pass or "Unit" in p_pass) else font_regular_sm
        c_pass.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

        ws_quote.merge_cells("J{0}:K{0}".format(s_idx))
        c_b = ws_quote.cell(row=s_idx, column=10, value=p_bill)
        c_b.font = font_bold_sm
        c_b.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for c in range(1, 12):
            cell = ws_quote.cell(row=s_idx, column=c)
            cell.border = box_border
            if fill_r.fill_type:
                cell.fill = fill_r

    # Row 55: Navigation Banner to Sheet 2
    ws_quote.merge_cells("A55:K55")
    ws_quote["A55"] = "📊 5.0 OPERATIONAL DISPATCH SCHEDULE: SEE SHEET 2 'Service_Gantt_Schedule' FOR MASTER 35-DAY TIMELINE & SHIFT ALLOCATION"
    ws_quote["A55"].font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    ws_quote["A55"].fill = fill_header
    ws_quote["A55"].alignment = align_center
    ws_quote.row_dimensions[55].height = 20
    for c in range(1, 12):
        ws_quote.cell(row=55, column=c).border = box_border

    # Page Break 2 (After Row 55 -> Ends Page 2)
    ws_quote.row_breaks.append(Break(id=55))

    # -------------------------------------------------------------
    # PAGE 3: SECTION 4.0 (4-Point Photographic Inspection)
    # -------------------------------------------------------------
    ws_quote.row_dimensions[56].height = 6
    ws_quote.merge_cells("A57:K57")
    ws_quote["A57"] = "4.0 LOBBY FLOOR CONDITION: 4-POINT PHOTO INSPECTION"
    ws_quote["A57"].font = font_section
    ws_quote.row_dimensions[57].height = 20

    ws_quote.merge_cells("A58:K58")
    ws_quote["A58"] = "Photographic documentation of floor conditions, wear patterns, and recommended services across portfolio lobbies."
    ws_quote["A58"].font = Font(name="Calibri", size=8.5, italic=True, color="475569")
    ws_quote["A58"].alignment = align_left
    ws_quote.row_dimensions[58].height = 14

    # Photo 1 & 2 Header (Row 59)
    ws_quote.row_dimensions[59].height = 18
    ws_quote.merge_cells("A59:E59")
    ws_quote["A59"] = "PHOTO 1: PLAZA HOTEL BUILDING (303 MAIN ST)"
    ws_quote["A59"].font = font_header
    ws_quote["A59"].fill = fill_header_opt
    ws_quote["A59"].alignment = align_center

    ws_quote.cell(row=59, column=6).value = "" # spacer

    ws_quote.merge_cells("G59:K59")
    ws_quote["G59"] = "PHOTO 2: CHASE BANK BUILDING (420 THROCKMORTON ST)"
    ws_quote["G59"].font = font_header
    ws_quote["G59"].fill = fill_header_opt
    ws_quote["G59"].alignment = align_center

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=59, column=c).border = box_border
            ws_quote.cell(row=59, column=c).fill = fill_header_opt

    # Image Viewport 1 & 2 (Row 60)
    ws_quote.row_dimensions[60].height = 150
    ws_quote.merge_cells("A60:E60")
    ws_quote.merge_cells("G60:K60")
    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=60, column=c).border = box_border
            ws_quote.cell(row=60, column=c).fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    img_plaza = OpenpyxlImage(img_paths["plaza"])
    img_plaza.width = 160
    img_plaza.height = 195
    ws_quote.add_image(img_plaza, "B60")

    img_chase = OpenpyxlImage(img_paths["chase"])
    img_chase.width = 160
    img_chase.height = 195
    ws_quote.add_image(img_chase, "H60")

    # Substrate (Row 61)
    ws_quote.row_dimensions[61].height = 16
    ws_quote.merge_cells("A61:E61")
    ws_quote["A61"] = "Primary floor system: Commercial Quarry Tile"
    ws_quote["A61"].font = font_bold_sm
    ws_quote["A61"].fill = fill_zebra

    ws_quote.merge_cells("G61:K61")
    ws_quote["G61"] = "Primary floor system: Polished Terrazzo"
    ws_quote["G61"].font = font_bold_sm
    ws_quote["G61"].fill = fill_zebra

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=61, column=c).border = box_border
            ws_quote.cell(row=61, column=c).fill = fill_zebra

    # Floor Issue (Row 62)
    ws_quote.row_dimensions[62].height = 16
    ws_quote.merge_cells("A62:E62")
    ws_quote["A62"] = "Floor Issue: Chemical Bleach Stains & Deep Grout Soil Buildup"
    ws_quote["A62"].font = font_alert

    ws_quote.merge_cells("G62:K62")
    ws_quote["G62"] = "Floor Issue: Heavy Traffic Wear Pattern & Loss of Floor Shine"
    ws_quote["G62"].font = font_alert

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=62, column=c).border = box_border

    # Condition Observed (Row 63)
    ws_quote.row_dimensions[63].height = 32
    ws_quote.merge_cells("A63:E63")
    ws_quote["A63"] = "Condition Observed: Harsh cleaner or bleach stripped the tile finish (white patch); grease soaked into unsealed tile; heavy dirt along grout lines."
    ws_quote["A63"].font = font_regular_sm
    ws_quote["A63"].alignment = align_desc_wrap

    ws_quote.merge_cells("G63:K63")
    ws_quote["G63"] = "Condition Observed: Heavy foot traffic wore away surface shine along center walkway; fine scratches now trap dirt and leave floors looking dull."
    ws_quote["G63"].font = font_regular_sm
    ws_quote["G63"].alignment = align_desc_wrap

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=63, column=c).border = box_border

    # Recommended Service (Row 64)
    ws_quote.row_dimensions[64].height = 30
    ws_quote.merge_cells("A64:E64")
    ws_quote["A64"] = "Recommended Service: Deep machine scrub to lift embedded soil (Phase 1) + Option A: BulletProof® Sealer to protect tile and grout from stains."
    ws_quote["A64"].font = font_bold_sm
    ws_quote["A64"].fill = fill_green
    ws_quote["A64"].alignment = align_desc_wrap

    ws_quote.merge_cells("G64:K64")
    ws_quote["G64"] = "Recommended Service: Regular quarterly machine scrub ($0.405/SF) to remove gritty soil and keep floors clean; deep gouges require Option C: Diamond Restoration."
    ws_quote["G64"].font = font_bold_sm
    ws_quote["G64"].fill = fill_green
    ws_quote["G64"].alignment = align_desc_wrap

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=64, column=c).border = box_border
            ws_quote.cell(row=64, column=c).fill = fill_green

    # Photo 3 & 4 (Rows 66–71)
    ws_quote.row_dimensions[65].height = 6

    ws_quote.row_dimensions[66].height = 18
    ws_quote.merge_cells("A66:E66")
    ws_quote["A66"] = "PHOTO 3: PETROLEUM BUILDING (210 W 6TH ST)"
    ws_quote["A66"].font = font_header
    ws_quote["A66"].fill = fill_header_opt
    ws_quote["A66"].alignment = align_center

    ws_quote.merge_cells("G66:K66")
    ws_quote["G66"] = "PHOTO 4: VIRTUOSO BUILDING (505 MAIN ST / SCHWARZ BLDG)"
    ws_quote["G66"].font = font_header
    ws_quote["G66"].fill = fill_header_opt
    ws_quote["G66"].alignment = align_center

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=66, column=c).border = box_border
            ws_quote.cell(row=66, column=c).fill = fill_header_opt

    ws_quote.row_dimensions[67].height = 150
    ws_quote.merge_cells("A67:E67")
    ws_quote.merge_cells("G67:K67")
    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=67, column=c).border = box_border
            ws_quote.cell(row=67, column=c).fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")

    img_pet = OpenpyxlImage(img_paths["petroleum"])
    img_pet.width = 160
    img_pet.height = 195
    ws_quote.add_image(img_pet, "B67")

    img_sch = OpenpyxlImage(img_paths["schwarz"])
    img_sch.width = 160
    img_sch.height = 195
    ws_quote.add_image(img_sch, "H67")

    ws_quote.row_dimensions[68].height = 16
    ws_quote.merge_cells("A68:E68")
    ws_quote["A68"] = "Primary floor system: Polished Marble Tile"
    ws_quote["A68"].font = font_bold_sm
    ws_quote["A68"].fill = fill_zebra

    ws_quote.merge_cells("G68:K68")
    ws_quote["G68"] = "Primary floor system: Polished Marble Tile"
    ws_quote["G68"].font = font_bold_sm
    ws_quote["G68"].fill = fill_zebra

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=68, column=c).border = box_border
            ws_quote.cell(row=68, column=c).fill = fill_zebra

    ws_quote.row_dimensions[69].height = 16
    ws_quote.merge_cells("A69:E69")
    ws_quote["A69"] = "Floor Issue: Acid Stain Marks & Chemical Etch Rings"
    ws_quote["A69"].font = font_alert

    ws_quote.merge_cells("G69:K69")
    ws_quote["G69"] = "Floor Issue: Deep Gouges & Scratches from Moving Equipment"
    ws_quote["G69"].font = font_alert

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=69, column=c).border = box_border

    ws_quote.row_dimensions[70].height = 32
    ws_quote.merge_cells("A70:E70")
    ws_quote["A70"] = "Condition Observed: Acidic spills (coffee, citrus drinks, sanitizer drips) ate into polished marble, leaving dull, rough white spots that mopping cannot clean."
    ws_quote["A70"].font = font_regular_sm
    ws_quote["A70"].alignment = align_desc_wrap

    ws_quote.merge_cells("G70:K70")
    ws_quote["G70"] = "Condition Observed: Heavy equipment, freight carts, or unpadded furniture dragged across floor cut deep white scratches into the black marble."
    ws_quote["G70"].font = font_regular_sm
    ws_quote["G70"].alignment = align_desc_wrap

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=70, column=c).border = box_border

    ws_quote.row_dimensions[71].height = 30
    ws_quote.merge_cells("A71:E71")
    ws_quote["A71"] = "Recommended Service: Diamond polishing to smooth out etch marks and restore shine + Option A: BulletProof® Sealer to protect against stains."
    ws_quote["A71"].font = font_bold_sm
    ws_quote["A71"].fill = fill_green
    ws_quote["A71"].alignment = align_desc_wrap

    ws_quote.merge_cells("G71:K71")
    ws_quote["G71"] = "Recommended Service: Option C: Diamond Restoration ($1.9875/SF) to grind out deep scratch valleys and polish back to a high-gloss finish."
    ws_quote["G71"].font = font_bold_sm
    ws_quote["G71"].fill = fill_green
    ws_quote["G71"].alignment = align_desc_wrap

    for c in range(1, 12):
        if c != 6:
            ws_quote.cell(row=71, column=c).border = box_border
            ws_quote.cell(row=71, column=c).fill = fill_green

    # Column Widths on Sheet 1 (Strictly 11 Columns, perfectly balanced for 100% full scale print)
    col_widths_s1 = {
        1: 5.5,   # Item
        2: 25.0,  # Building Name
        3: 22.0,  # Address
        4: 24.0,  # Primary Floor Substrate
        5: 14.0,  # Cleanable Area (SQFT)
        6: 13.0,  # Phase 1 Rate
        7: 18.0,  # Phase 1 deep clean
        8: 18.0,  # Quarterly Rate & Option B
        9: 16.0,  # Track A: Quarterly
        10: 14.0, # Semi-Annual Rate
        11: 19.5  # Track B: Semi-Annual & Option D
    }
    for col_idx, width in col_widths_s1.items():
        ws_quote.column_dimensions[get_column_letter(col_idx)].width = width

    # PRINT SETUP FOR SHEET 1
    # Mandated by User: Repeat entire page header down to Data Integrity Note ($1:$8)
    ws_quote.print_title_rows = '$1:$8'
    ws_quote.page_setup.orientation = ws_quote.ORIENTATION_LANDSCAPE
    ws_quote.page_setup.paperSize = ws_quote.PAPERSIZE_LETTER
    ws_quote.sheet_properties.pageSetUpPr.fitToPage = True
    ws_quote.page_setup.fitToWidth = 1
    ws_quote.page_setup.fitToHeight = 0
    ws_quote.page_margins.left = 0.35
    ws_quote.page_margins.right = 0.35
    ws_quote.page_margins.top = 0.35
    ws_quote.page_margins.bottom = 0.35

    # =============================================================
    # SHEET 2: Service_Gantt_Schedule (Dedicated 1-Page Landscape Attachment)
    # =============================================================
    ws_gantt = wb.create_sheet(title="Service_Gantt_Schedule")
    ws_gantt.views.sheetView[0].showGridLines = True

    ws_gantt.merge_cells("A1:AN1")
    ws_gantt["A1"] = "HWB CLEANING SERVICES LLC"
    ws_gantt["A1"].font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws_gantt["A1"].fill = fill_title
    ws_gantt["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_gantt.row_dimensions[1].height = 22

    ws_gantt.merge_cells("A2:AN2")
    ws_gantt["A2"] = "OPERATIONAL DISPATCH GANTT SCHEDULE — BOSANNA LLC / 11-BUILDING COMMERCIAL PORTFOLIO"
    ws_gantt["A2"].font = Font(name="Calibri", size=9.5, bold=True, color="93C5FD")
    ws_gantt["A2"].fill = fill_title
    ws_gantt["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_gantt.row_dimensions[2].height = 16

    ws_gantt.row_dimensions[3].height = 18
    ws_gantt.merge_cells("A3:B3")
    ws_gantt["A3"] = "Dispatch Origin: Arlington, TX"
    ws_gantt["A3"].font = font_bold
    ws_gantt["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws_gantt.merge_cells("C3:D3")
    ws_gantt["C3"] = "Target Start Date: Sep 01, 2026"
    ws_gantt["C3"].font = font_bold
    ws_gantt["C3"].alignment = Alignment(horizontal="center", vertical="center")

    ws_gantt.merge_cells("E3:I3")
    ws_gantt["E3"] = "Assigned Crew: 2 Technicians"
    ws_gantt["E3"].font = font_regular
    ws_gantt["E3"].alignment = Alignment(horizontal="center", vertical="center")

    ws_gantt.merge_cells("J3:Q3")
    ws_gantt["J3"] = "Mobilization: 10 Dedicated Shifts"
    ws_gantt["J3"].font = font_regular
    ws_gantt["J3"].alignment = Alignment(horizontal="center", vertical="center")

    ws_gantt.merge_cells("R3:AN3")
    ws_gantt["R3"] = "⏰ 100% Flexible Shift Timing (Client Choice: Zero Tenant Disruption)"
    ws_gantt["P3"].font = Font(name="Calibri", size=8.5, bold=True, color="16A34A")
    ws_gantt["P3"].alignment = Alignment(horizontal="center", vertical="center")

    for c in range(1, 41):
        ws_gantt.cell(row=3, column=c).border = box_border

    ws_gantt.row_dimensions[4].height = 18
    ws_gantt.merge_cells("A4:E4")
    ws_gantt["A4"] = "LEGEND & COLOR CODES:"
    ws_gantt["A4"].font = font_bold
    ws_gantt["A4"].fill = fill_subtle
    ws_gantt["A4"].alignment = align_center

    legend_items_s2 = [
        ("F4:L4", "DP: Phase 1 Deep Clean", fill_p1_deep_clean),
        ("M4:S4", "SL: Modular Sealer", fill_p2_sealer),
        ("T4:Z4", "QT: Quarterly Care", fill_p3_quarterly),
        ("AA4:AG4", "ST: Stair Detailing", fill_opt_stairs),
        ("AH4:AN4", "WKND: Building Offline", PatternFill(start_color="CBD5E1", end_color="CBD5E1", fill_type="solid"))
    ]
    for crange, ltext, lfill in legend_items_s2:
        ws_gantt.merge_cells(crange)
        c_top = ws_gantt[crange.split(":")[0]]
        c_top.value = ltext
        c_top.font = font_badge
        c_top.fill = lfill
        c_top.alignment = align_center
        c_start, _, c_end, _ = range_boundaries(crange)
        for col_i in range(c_start, c_end + 1):
            ws_gantt.cell(row=4, column=col_i).border = box_border
            ws_gantt.cell(row=4, column=col_i).fill = lfill

    for c in range(1, 6):
        ws_gantt.cell(row=4, column=c).border = box_border

    ws_gantt.row_dimensions[5].height = 18
    ws_gantt.row_dimensions[6].height = 15
    ws_gantt.row_dimensions[7].height = 15

    headers_left_s2 = [
        ("A", "Item"),
        ("B", "Building Name"),
        ("C", "Total SQF"),
        ("D", "Shift Window"),
        ("E", "Crew Hrs")
    ]
    for col_let, htext in headers_left_s2:
        ws_gantt.merge_cells(f"{col_let}5:{col_let}7")
        c_h = ws_gantt[f"{col_let}5"]
        c_h.value = htext
        c_h.font = font_header
        c_h.fill = fill_header_opt
        c_h.alignment = align_header_wrap
        for r in range(5, 8):
            ws_gantt[f"{col_let}{r}"].border = box_border
            ws_gantt[f"{col_let}{r}"].fill = fill_header_opt

    weeks_s2 = [
        ("F5:L5", 6, 12, "Week 1 (Aug 31–Sep 06)"),
        ("M5:S5", 13, 19, "Week 2 (Sep 07–Sep 13)"),
        ("T5:Z5", 20, 26, "Week 3 (Sep 14–Sep 20)"),
        ("AA5:AG5", 27, 33, "Week 4 (Sep 21–Sep 27)"),
        ("AH5:AN5", 34, 40, "Week 5 (Sep 28–Oct 04)")
    ]
    for wrange, start_c, end_c, wtitle in weeks_s2:
        ws_gantt.merge_cells(wrange)
        c_w = ws_gantt[wrange.split(":")[0]]
        c_w.value = wtitle
        c_w.font = font_header
        c_w.fill = fill_header
        c_w.alignment = align_center
        for c in range(start_c, end_c + 1):
            ws_gantt.cell(row=5, column=c).border = box_border
            ws_gantt.cell(row=5, column=c).fill = fill_header

    days_dates = [31] + list(range(1, 31)) + list(range(1, 5))
    days_names = ['Mo', 'Tu', 'We', 'Th', 'Fr', 'Sa', 'Su'] * 5

    for d_idx, d_num in enumerate(days_dates):
        col = 6 + d_idx
        c_num = ws_gantt.cell(row=6, column=col, value=d_num)
        c_num.font = Font(name="Calibri", size=8.5, bold=True, color="FFFFFF")
        c_num.fill = fill_title
        c_num.alignment = align_center
        c_num.border = box_border

    for d_idx, d_name in enumerate(days_names):
        col = 6 + d_idx
        c_nam = ws_gantt.cell(row=7, column=col, value=d_name)
        is_w = d_name in ['Sa', 'Su']
        c_nam.fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid") if is_w else PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        c_nam.font = Font(name="Calibri", size=8, bold=True, color="FFFFFF" if is_w else "94A3B8")
        c_nam.alignment = align_center
        c_nam.border = box_border

    schedule_data_s2 = [
        (1, "Chase Bank Building", 2126, "Dedicated Night 1 (3.5h)", 3.5, [7], [11], [21], []),
        (2, "The Westbrook", 1520, "Dedicated Night 2 (3.0h)", 3.0, [8], [11], [22], []),
        (3, "The Carnegie", 1175, "Dedicated Night 3 (2.8h)", 2.8, [9], [11], [23], []),
        (4, "The Cassidy", 1134, "Dedicated Night 4 (2.6h)", 2.6, [10], [11], [24], []),
        (5, "Burk Burnett Building", 1108, "Dedicated Night 5 (2.6h)", 2.6, [13], [18], [27], [28]),
        (6, "Sanger Lofts", 1071, "Dedicated Night 6 (2.5h)", 2.5, [14], [18], [28], []),
        (7, "Petroleum Building", 860, "Dedicated Night 7 (2.2h)", 2.2, [15], [18], [29], []),
        (8, "The Commerce Building", 840, "Dedicated Night 8 (2.2h)", 2.2, [16], [18], [30], []),
        (9, "Virtuoso Building", 560, "Boutique Night 9 (1.8h)", 1.8, [17], [18], [31], []),
        (10, "Knights of Pythias Hall", 414, "Boutique Night 9 (1.6h)", 1.6, [17], [18], [31], []),
        (11, "Plaza Hotel Building", 345, "Boutique Night 10 (1.5h)", 1.5, [20], [21], [34], [])
    ]

    for b_idx, item in enumerate(schedule_data_s2, start=8):
        num, name, sf, window, crew_time, p1_cols, seal_cols, rec_cols, stair_cols = item
        ws_gantt.row_dimensions[b_idx].height = 16

        ws_gantt.cell(row=b_idx, column=1, value=num).alignment = align_center
        ws_gantt.cell(row=b_idx, column=2, value=name).alignment = align_left
        
        c_sf = ws_gantt.cell(row=b_idx, column=3, value=sf)
        c_sf.alignment = align_center
        c_sf.number_format = "#,##0"

        ws_gantt.cell(row=b_idx, column=4, value=window).alignment = align_left

        c_hrs = ws_gantt.cell(row=b_idx, column=5, value=crew_time)
        c_hrs.alignment = align_center
        c_hrs.number_format = '0.0 "hrs"'

        for c in range(1, 6):
            cell = ws_gantt.cell(row=b_idx, column=c)
            cell.font = font_bold if c in [2, 3] else font_regular
            cell.border = box_border
            if b_idx % 2 == 1:
                cell.fill = fill_zebra

        for day_i in range(1, 36):
            col = day_i + 5
            cell = ws_gantt.cell(row=b_idx, column=col)
            cell.border = box_border

            if col in p1_cols:
                cell.fill = fill_p1_deep_clean
                cell.value = "DP"
                cell.font = font_badge
                cell.alignment = align_center
            elif col in seal_cols:
                cell.fill = fill_p2_sealer
                cell.value = "SL"
                cell.font = font_badge
                cell.alignment = align_center
            elif col in rec_cols:
                cell.fill = fill_p3_quarterly
                cell.value = "QT"
                cell.font = font_badge
                cell.alignment = align_center
            elif col in stair_cols:
                cell.fill = fill_opt_stairs
                cell.value = "ST"
                cell.font = font_badge
                cell.alignment = align_center
            elif (col - 6) % 7 in [5, 6]:
                cell.fill = fill_weekend

    tot_r_s2 = 19
    ws_gantt.row_dimensions[tot_r_s2].height = 18
    ws_gantt.merge_cells("A19:B19")
    ws_gantt["A19"] = "TOTAL PORTFOLIO DISPATCH"
    ws_gantt["A19"].font = font_bold
    ws_gantt["A19"].alignment = align_center

    c_tsf = ws_gantt.cell(row=tot_r_s2, column=3, value="=SUM(C8:C18)")
    c_tsf.font = font_bold
    c_tsf.number_format = '#,##0 " SQF"'
    c_tsf.alignment = align_center

    ws_gantt.cell(row=tot_r_s2, column=4, value="10 Dedicated Shifts").alignment = align_center
    c_thrs = ws_gantt.cell(row=tot_r_s2, column=5, value="=SUM(E8:E18)")
    c_thrs.font = font_bold
    c_thrs.number_format = '0.0 "hrs total"'
    c_thrs.alignment = align_center

    for c in range(1, 6):
        ws_gantt.cell(row=tot_r_s2, column=c).border = total_border
        ws_gantt.cell(row=tot_r_s2, column=c).fill = fill_subtle

    for c in range(6, 41):
        cell = ws_gantt.cell(row=tot_r_s2, column=c)
        cell.border = total_border
        cell.fill = fill_subtle

    ws_gantt.row_dimensions[20].height = 6
    ws_gantt.merge_cells("A21:AN21")
    ws_gantt["A21"] = "*Note: Shift timing and scheduled calendar dates are 100% flexible based on client preference. HWB will adapt dispatch hours (earlier start, late night, or weekend windows) to accommodate building security and tenant schedules at no additional cost."
    ws_gantt["A21"].font = Font(name="Calibri", size=8, italic=True, color="475569")
    ws_gantt["A21"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws_gantt.row_dimensions[21].height = 24

    ws_gantt.column_dimensions['A'].width = 6.0
    ws_gantt.column_dimensions['B'].width = 26.0
    ws_gantt.column_dimensions['C'].width = 14.0
    ws_gantt.column_dimensions['D'].width = 24.0
    ws_gantt.column_dimensions['E'].width = 15.0
    for col_idx in range(6, 41):
        ws_gantt.column_dimensions[get_column_letter(col_idx)].width = 4.5

    ws_gantt.print_title_rows = '$1:$4'
    ws_gantt.page_setup.orientation = ws_gantt.ORIENTATION_LANDSCAPE
    ws_gantt.page_setup.paperSize = ws_gantt.PAPERSIZE_LETTER
    ws_gantt.sheet_properties.pageSetUpPr.fitToPage = True
    ws_gantt.page_setup.fitToWidth = 1
    ws_gantt.page_setup.fitToHeight = 1
    ws_gantt.page_margins.left = 0.25
    ws_gantt.page_margins.right = 0.25
    ws_gantt.page_margins.top = 0.35
    ws_gantt.page_margins.bottom = 0.35
    ws_gantt.print_area = 'A1:AN21'

    # =============================================================
    # SHEET 3: Cost_Underwriting_Audit
    # =============================================================
    ws_cost = wb.create_sheet(title="Cost_Underwriting_Audit")
    ws_cost.views.sheetView[0].showGridLines = True

    ws_cost.merge_cells("A1:K1")
    ws_cost["A1"] = "HWB CLEANING SERVICES LLC"
    ws_cost["A1"].font = font_title
    ws_cost["A1"].fill = fill_title
    ws_cost["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_cost.row_dimensions[1].height = 24

    ws_cost.merge_cells("A2:K2")
    ws_cost["A2"] = "INTERNAL OPERATIONAL UNDERWRITING AUDIT — BOSANNA LLC DISPATCH BASIS"
    ws_cost["A2"].font = Font(name="Calibri", size=9.5, bold=True, color="93C5FD")
    ws_cost["A2"].fill = fill_title
    ws_cost["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_cost.row_dimensions[2].height = 16

    cost_headers = [
        "Item",
        "Building Name",
        "Area (SF)",
        "Crew Hours\n(2 Techs)",
        "Labor Cost\n($57.50/hr)",
        "Chemicals\n& Supplies",
        "Transit & Equip",
        "Supervision\n& Admin",
        "Total Cost\nBasis",
        "Quarterly Quote\n($0.405/SF)",
        "Gross Profit\nMargin %"
    ]
    ws_cost.row_dimensions[4].height = 36
    for col_num, h_text in enumerate(cost_headers, 1):
        cell = ws_cost.cell(row=4, column=col_num, value=h_text)
        cell.font = font_header
        cell.fill = fill_header_opt
        cell.alignment = align_header_wrap
        cell.border = box_border

    cost_data = [
        (1, "Chase Bank Building", 2126, 3.50, 45.0, 95.0, 60.0),
        (2, "The Westbrook", 1520, 3.00, 35.0, 95.0, 50.0),
        (3, "The Carnegie", 1175, 2.75, 28.0, 95.0, 45.0),
        (4, "The Cassidy", 1134, 2.60, 26.0, 95.0, 45.0),
        (5, "Burk Burnett Building", 1108, 2.60, 26.0, 95.0, 45.0),
        (6, "Sanger Lofts", 1071, 2.50, 25.0, 95.0, 40.0),
        (7, "Petroleum Building", 860, 2.20, 20.0, 95.0, 35.0),
        (8, "The Commerce Building", 840, 2.20, 20.0, 95.0, 35.0),
        (9, "Virtuoso Building", 560, 1.80, 15.0, 95.0, 30.0),
        (10, "Knights of Pythias Hall", 414, 1.60, 12.0, 95.0, 25.0),
        (11, "Plaza Hotel Building", 345, 1.50, 10.0, 95.0, 25.0)
    ]

    for idx, item in enumerate(cost_data, start=5):
        num, name, sf, hrs, chem, tr_eq, misc = item
        ws_cost.row_dimensions[idx].height = 16
        fill_r = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)

        ws_cost.cell(row=idx, column=1, value=num).alignment = align_center
        ws_cost.cell(row=idx, column=2, value=name).alignment = align_left
        
        c_sf = ws_cost.cell(row=idx, column=3, value=sf)
        c_sf.number_format = "#,##0"
        c_sf.alignment = align_center

        c_hrs = ws_cost.cell(row=idx, column=4, value=hrs)
        c_hrs.number_format = "0.00"
        c_hrs.alignment = align_right

        c_lab = ws_cost.cell(row=idx, column=5, value="=D{0}*57.50".format(idx))
        c_lab.number_format = "$#,##0.00"
        c_lab.alignment = align_right

        c_ch = ws_cost.cell(row=idx, column=6, value=chem)
        c_ch.number_format = "$#,##0.00"
        c_ch.alignment = align_right

        c_tr = ws_cost.cell(row=idx, column=7, value=tr_eq)
        c_tr.number_format = "$#,##0.00"
        c_tr.alignment = align_right

        c_mi = ws_cost.cell(row=idx, column=8, value=misc)
        c_mi.number_format = "$#,##0.00"
        c_mi.alignment = align_right

        c_tot = ws_cost.cell(row=idx, column=9, value="=SUM(E{0}:H{0})".format(idx))
        c_tot.number_format = "$#,##0.00"
        c_tot.alignment = align_right

        c_qp = ws_cost.cell(row=idx, column=10, value="=MAX(C{0},600)*0.405".format(idx))
        c_qp.number_format = "$#,##0.00"
        c_qp.alignment = align_right

        c_gm = ws_cost.cell(row=idx, column=11, value="=(J{0}-I{0})/J{0}".format(idx))
        c_gm.number_format = "0.0%"
        c_gm.alignment = align_right

        for c in range(1, 12):
            cell = ws_cost.cell(row=idx, column=c)
            cell.font = font_bold if c in [2, 9, 10, 11] else font_regular
            cell.border = box_border
            if fill_r.fill_type:
                cell.fill = fill_r

    tot_c_row = 16
    ws_cost.row_dimensions[tot_c_row].height = 28
    ws_cost.merge_cells("A{0}:B{0}".format(tot_c_row))
    ws_cost["A{0}".format(tot_c_row)] = "TOTAL PER QUARTERLY PASS\n(10 Dedicated Night Shifts)"
    ws_cost["A{0}".format(tot_c_row)].font = font_bold
    ws_cost["A{0}".format(tot_c_row)].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    c_tot_csf = ws_cost.cell(row=tot_c_row, column=3, value="=SUM(C5:C15)")
    c_tot_csf.number_format = '#,##0 " SQF"'
    c_tot_csf.alignment = align_center

    ws_cost.cell(row=tot_c_row, column=4, value="=SUM(D5:D15)").number_format = "0.00"
    ws_cost.cell(row=tot_c_row, column=5, value="=SUM(E5:E15)").number_format = "$#,##0.00"
    ws_cost.cell(row=tot_c_row, column=6, value="=SUM(F5:F15)").number_format = "$#,##0.00"
    ws_cost.cell(row=tot_c_row, column=7, value="=SUM(G5:G15)").number_format = "$#,##0.00"
    ws_cost.cell(row=tot_c_row, column=8, value="=SUM(H5:H15)").number_format = "$#,##0.00"
    ws_cost.cell(row=tot_c_row, column=9, value="=SUM(I5:I15)").number_format = "$#,##0.00"
    ws_cost.cell(row=tot_c_row, column=10, value="=SUM(J5:J15)").number_format = "$#,##0.00"
    ws_cost.cell(row=tot_c_row, column=11, value="=(J{0}-I{0})/J{0}".format(tot_c_row)).number_format = "0.0%"

    for c in range(1, 12):
        cell = ws_cost.cell(row=tot_c_row, column=c)
        cell.font = font_bold
        cell.fill = fill_subtle
        cell.border = total_border
        if c >= 4:
            cell.alignment = align_right

    cost_widths = {1: 6, 2: 24, 3: 16, 4: 16, 5: 16, 6: 14, 7: 18, 8: 15, 9: 16, 10: 16, 11: 15}
    for col_idx, width in cost_widths.items():
        ws_cost.column_dimensions[get_column_letter(col_idx)].width = width

    ws_cost.print_title_rows = '$1:$4'
    ws_cost.page_setup.orientation = ws_cost.ORIENTATION_LANDSCAPE
    ws_cost.page_setup.paperSize = ws_cost.PAPERSIZE_LETTER
    ws_cost.sheet_properties.pageSetUpPr.fitToPage = True
    ws_cost.page_setup.fitToWidth = 1
    ws_cost.page_setup.fitToHeight = 1
    ws_cost.page_margins.left = 0.25
    ws_cost.page_margins.right = 0.25
    ws_cost.page_margins.top = 0.35
    ws_cost.page_margins.bottom = 0.35

    header_institution = "&B&10HWB CLEANING SERVICES LLC"
    footer_page_no = "Page &P of &N"

    for ws in [ws_quote, ws_gantt, ws_cost]:
        ws.oddHeader.center.text = header_institution
        ws.oddHeader.center.size = 10
        ws.oddHeader.center.font = "Calibri,Bold"
        ws.evenHeader.center.text = header_institution
        ws.oddFooter.right.text = footer_page_no
        ws.evenFooter.right.text = footer_page_no

    save_destinations = [
        "HWB-COMPANY/HWB-QUOTES/BOSANNA-SUNDANCE-QUOTE-gantt.xlsx",
        "HWB-COMPANY/HWB-QUOTES/BOSANNA-SUNDANCE-QUOTE.xlsx",
        "HWB-COMPANY/HWB-QUOTES/BOSANNA-QUOTE-gantt.xlsx",
        "HWB-COMPANY/HWB-QUOTES/SUNDANCE-QUOTE-gantt.xlsx"
    ]
    for dest in save_destinations:
        wb.save(dest)
        print(f"SUCCESS: Saved Option 1 proposal workbook to {dest}")

if __name__ == "__main__":
    generate_option1_workbook()
