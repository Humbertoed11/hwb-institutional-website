import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, range_boundaries

def create_gc_bid_excel():
    target_path = "HWB-COMPANY/HWB-ESTIMATING/TEMPLATES/HWB-CSI-017423-Commercial-Final-Cleaning-Bid-Form.xlsx"
    wb = openpyxl.Workbook()

    # Styling definitions
    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=9.5, bold=True, color="93C5FD")
    font_section = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=9, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=9, color="334155")
    font_sm = Font(name="Calibri", size=8, color="64748B")
    font_total = Font(name="Calibri", size=11, bold=True, color="1E3A8A")

    fill_dark = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_blue = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_subtle = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_green = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    fill_alert = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    thin_side = Side(style="thin", color="CBD5E1")
    box_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
    total_border = Border(top=Side(style="thin", color="0F172A"), bottom=Side(style="double", color="0F172A"), left=thin_side, right=thin_side)

    # -------------------------------------------------------------
    # SHEET 1: Subcontract_Bid_Form
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Subcontract_Bid_Form"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner (Rows 1 & 2)
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "HWB CLEANING SERVICES LLC — COMMERCIAL CONSTRUCTION DIVISION"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_dark
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[1].height = 26

    ws1.merge_cells("A2:G2")
    ws1["A2"] = "CSI MASTERFORMAT® 01 74 23: FINAL CLEANING SUBCONTRACT BID PROPOSAL"
    ws1["A2"].font = font_sub
    ws1["A2"].fill = fill_dark
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[2].height = 16

    ws1.row_dimensions[3].height = 6

    # Metadata Box (Rows 4 to 7)
    meta = [
        ("General Contractor:", "Novel Builders LLC", "Bid Date:", "September 04, 2026"),
        ("Project Name:", "Stacked Industrial", "Bid Ref ID:", "HWB-GC-2026-001"),
        ("Project Location:", "Fort Worth / DFW Metroplex, TX", "Addenda Ack:", "Addenda #1 through #3 (All Current)"),
        ("Estimator Contact:", "Austin Addis", "Operations Base:", "Arlington, TX (Dedicated DFW Mobilization)")
    ]

    for idx, (l1, v1, l2, v2) in enumerate(meta, start=4):
        ws1.row_dimensions[idx].height = 18
        ws1.merge_cells(f"A{idx}:B{idx}")
        ws1[f"A{idx}"] = l1
        ws1[f"A{idx}"].font = font_bold
        ws1[f"A{idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws1.merge_cells(f"C{idx}:D{idx}")
        ws1[f"C{idx}"] = v1
        ws1[f"C{idx}"].font = font_regular
        ws1[f"C{idx}"].alignment = Alignment(horizontal="left", vertical="center")

        ws1[f"E{idx}"] = l2
        ws1[f"E{idx}"].font = font_bold
        ws1[f"E{idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws1.merge_cells(f"F{idx}:G{idx}")
        ws1[f"F{idx}"] = v2
        ws1[f"F{idx}"].font = font_regular
        ws1[f"F{idx}"].alignment = Alignment(horizontal="left", vertical="center")

        for c in range(1, 8):
            cell = ws1.cell(row=idx, column=c)
            cell.border = box_border
            if c in [1, 2, 5]:
                cell.fill = fill_zebra

    ws1.row_dimensions[8].height = 6

    # Section 1.0 Header
    ws1.merge_cells("A9:G9")
    ws1["A9"] = "1.0 CSI 01 74 23 ITEMIZED SCOPE OF WORK & SUBCONTRACT BID BREAKDOWN"
    ws1["A9"].font = font_section
    ws1["A9"].fill = fill_blue
    ws1["A9"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[9].height = 20

    # Table 1 Headers
    t1_headers = [
        ("A10", "Item", 6.0),
        ("B10:C10", "Phase / Scope Description", 38.0),
        ("D10", "Area / Units", 18.0),
        ("E10", "Rate Basis", 16.0),
        ("F10", "Total Investment", 20.0),
        ("G10", "Schedule Tier", 16.0)
    ]
    ws1.row_dimensions[10].height = 28
    for col_info in t1_headers:
        c_range = col_info[0]
        h_text = col_info[1]
        if ":" in c_range:
            ws1.merge_cells(c_range)
            top_c = ws1[c_range.split(":")[0]]
        else:
            top_c = ws1[c_range]
        top_c.value = h_text
        top_c.font = font_header
        top_c.fill = fill_dark
        top_c.alignment = Alignment(horizontal="center", vertical="center")

    for c in range(1, 8):
        ws1.cell(row=10, column=c).border = box_border
        ws1.cell(row=10, column=c).fill = fill_dark

    # Table 1 Rows
    bid_items = [
        ("1.1", "Phase 2: Master Final Clean (Primary Turnover)", "18,500 SF", 0.24, "=D11*0.24", "Base Scope"),
        ("1.2", "Glazing & Storefront Glass Detailing (In & Out)", "48 Panels", 18.00, "=48*18", "Base Scope"),
        ("1.3", "Restroom Deep De-Scale & Fixture Detailing", "14 Fixtures", 35.00, "=14*35", "Base Scope"),
        ("1.4", "Resilient Flooring Machine Scrub & Wax (VCT/LVT)", "3,200 SF", 0.28, "=3200*0.28", "Base Scope"),
        ("1.5", "Sealed / Polished Concrete Mechanical Auto-Scrub", "15,300 SF", 0.18, "=15300*0.18", "Base Scope"),
        ("1.6", "Phase 3: Pre-CO Touch-Up Clean (Fluff-and-Buff)", "18,500 SF", 0.06, "=D16*0.06", "Included Alternate")
    ]

    for idx, item in enumerate(bid_items, start=11):
        ws1.row_dimensions[idx].height = 20
        fill_row = fill_zebra if idx % 2 == 0 else PatternFill(fill_type=None)
        
        ws1.cell(row=idx, column=1, value=item[0]).alignment = Alignment(horizontal="center", vertical="center")
        
        ws1.merge_cells(f"B{idx}:C{idx}")
        ws1[f"B{idx}"] = item[1]
        ws1[f"B{idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        c_d = ws1.cell(row=idx, column=4)
        if "SF" in item[2]:
            val_num = int(item[2].replace(" SF", "").replace(",", ""))
            c_d.value = val_num
            c_d.number_format = '#,##0 " SF"'
        else:
            c_d.value = item[2]
        c_d.alignment = Alignment(horizontal="center", vertical="center")

        c_e = ws1.cell(row=idx, column=5, value=item[3])
        c_e.number_format = "$#,##0.00"
        c_e.alignment = Alignment(horizontal="right", vertical="center")

        c_f = ws1.cell(row=idx, column=6)
        if idx == 11:
            c_f.value = "=D11*E11"
        elif idx == 14:
            c_f.value = "=3200*E14"
        elif idx == 15:
            c_f.value = "=15300*E15"
        elif idx == 16:
            c_f.value = "=D16*E16"
        else:
            c_f.value = item[4]
        c_f.number_format = "$#,##0.00"
        c_f.alignment = Alignment(horizontal="right", vertical="center")

        ws1.cell(row=idx, column=7, value=item[5]).alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 8):
            cell = ws1.cell(row=idx, column=c)
            cell.font = font_bold if c in [2, 6] else font_regular
            cell.border = box_border
            if fill_row.fill_type:
                cell.fill = fill_row

    # Total Row
    ws1.row_dimensions[17].height = 24
    ws1.merge_cells("A17:E17")
    ws1["A17"] = "TOTAL BASE BID SUBCONTRACT INVESTMENT (Turnkey Fixed Sum)"
    ws1["A17"].font = font_bold
    ws1["A17"].alignment = Alignment(horizontal="center", vertical="center")

    c_tot = ws1.cell(row=17, column=6, value="=SUM(F11:F16)")
    c_tot.font = font_total
    c_tot.number_format = "$#,##0.00"
    c_tot.alignment = Alignment(horizontal="right", vertical="center")

    ws1.cell(row=17, column=7, value="CSI 01 74 23").alignment = Alignment(horizontal="center", vertical="center")

    for c in range(1, 8):
        cell = ws1.cell(row=17, column=c)
        cell.fill = fill_subtle
        cell.border = total_border

    ws1.row_dimensions[18].height = 6

    # Section 2.0 Inclusions & Exclusions
    ws1.merge_cells("A19:G19")
    ws1["A19"] = "2.0 SCOPE BOUNDARIES: POKA-YOKE INCLUSIONS & EXCLUSIONS"
    ws1["A19"].font = font_section
    ws1["A19"].fill = fill_blue
    ws1["A19"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[19].height = 20

    ws1.merge_cells("A20:C20")
    ws1["A20"] = "✓ SCOPE INCLUSIONS (By HWB Cleaning Services)"
    ws1["A20"].font = Font(name="Calibri", size=9, bold=True, color="15803D")
    ws1["A20"].fill = fill_green

    ws1.cell(row=20, column=4).value = ""

    ws1.merge_cells("E20:G20")
    ws1["E20"] = "✕ SCOPE EXCLUSIONS (By General Contractor / Others)"
    ws1["E20"].font = Font(name="Calibri", size=9, bold=True, color="B91C1C")
    ws1["E20"].fill = fill_alert

    inclusions = [
        "Sticker, protective film, and tape scraping on all glazing/frames.",
        "Storefront interior & exterior glass washed/squeegeed up to 25 ft.",
        "Machine rotary scrub and grout extraction across all tile/restrooms.",
        "Auto-scrubber mechanical wash on sealed/polished concrete.",
        "Acid wash and descaling of new plumbing porcelain & fixtures.",
        "Millwork, casework, and breakroom cabinetry interior/exterior wipe.",
        "All commercial chemicals, HEPA vacuums, scrubbers, and PPE."
    ]

    exclusions = [
        "On-site dumpster / disposal container (provided by GC on site).",
        "Active electrical power and pressurized water on site.",
        "Exterior glass exceeding 2 stories requiring boom/swing-stage.",
        "Hazardous material abatement (asbestos, lead, mold, biohazard).",
        "Pressure washing exterior parking lot / drive lanes.",
        "Repair of pre-existing scratch damage caused by prior trades.",
        "Trade re-cleaning after superintendent punch-walk ($45/hr T&M)."
    ]

    for idx, (inc, exc) in enumerate(zip(inclusions, exclusions), start=21):
        ws1.row_dimensions[idx].height = 18
        ws1.merge_cells(f"A{idx}:C{idx}")
        ws1[f"A{idx}"] = f"• {inc}"
        ws1[f"A{idx}"].font = font_regular
        ws1[f"A{idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws1.cell(row=idx, column=4).value = ""

        ws1.merge_cells(f"E{idx}:G{idx}")
        ws1[f"E{idx}"] = f"• {exc}"
        ws1[f"E{idx}"].font = font_regular
        ws1[f"E{idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        for c in range(1, 4):
            ws1.cell(row=idx, column=c).border = box_border
        for c in range(5, 8):
            ws1.cell(row=idx, column=c).border = box_border

    ws1.row_dimensions[28].height = 6

    # Section 3.0 Insurance & Terms
    ws1.merge_cells("A29:G29")
    ws1["A29"] = "3.0 COMMERCIAL INSURANCE, SAFETY (EMR) & CONTRACT TERMS"
    ws1["A29"].font = font_section
    ws1["A29"].fill = fill_blue
    ws1["A29"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws1.row_dimensions[29].height = 20

    ws1.merge_cells("A30:G30")
    ws1["A30"] = "COMPLIANCE: $2,000,000 General Liability | $1,000,000 Auto | $1,000,000 Workers' Comp | EMR 0.82 | OSHA 30 Certified"
    ws1["A30"].font = Font(name="Calibri", size=8.5, bold=True, color="166534")
    ws1["A30"].fill = fill_green
    ws1["A30"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[30].height = 18
    for c in range(1, 8):
        ws1.cell(row=30, column=c).border = box_border

    terms = [
        ("Billing & Retainage:", "Monthly progress billing (AIA G702/G703). 5–10% retainage released within 30 days of final clean turnover."),
        ("Validity & Mobilization:", "Proposal valid for 60 days. Mobilization requires 5 business days advance notice from Superintendent."),
        ("1-Click Prequal Vault:", "Download ACORD 25 COI, W-9, Safety Statement (IIPP), and References at https://mop.hwbcleaning.com/prequal")
    ]
    for idx, (t_lbl, t_val) in enumerate(terms, start=31):
        ws1.row_dimensions[idx].height = 18
        ws1.merge_cells(f"A{idx}:B{idx}")
        ws1[f"A{idx}"] = t_lbl
        ws1[f"A{idx}"].font = font_bold
        ws1[f"A{idx}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws1.merge_cells(f"C{idx}:G{idx}")
        ws1[f"C{idx}"] = t_val
        ws1[f"C{idx}"].font = font_regular
        ws1[f"C{idx}"].alignment = Alignment(horizontal="left", vertical="center")

        for c in range(1, 8):
            cell = ws1.cell(row=idx, column=c)
            cell.border = box_border
            if c in [1, 2]:
                cell.fill = fill_zebra

    # Signatures
    ws1.row_dimensions[34].height = 6
    ws1.merge_cells("A35:C35")
    ws1["A35"] = "SUBMITTED BY: HWB CLEANING SERVICES LLC"
    ws1["A35"].font = font_bold

    ws1.merge_cells("E35:G35")
    ws1["E35"] = "ACCEPTED BY: GENERAL CONTRACTOR"
    ws1["E35"].font = font_bold

    ws1.row_dimensions[36].height = 24
    ws1.merge_cells("A36:C36")
    ws1["A36"] = "Humberto Dominguez, CEO • (817) 600-6467"
    ws1["A36"].font = font_regular

    ws1.merge_cells("E36:G36")
    ws1["E36"] = "Authorized Signature: ___________________________________"
    ws1["E36"].font = font_regular

    # Column Widths on Sheet 1
    col_w = {
        1: 6.0,   # Item
        2: 24.0,  # Phase Desc A
        3: 16.0,  # Phase Desc B
        4: 18.0,  # Area / Units
        5: 16.0,  # Rate Basis
        6: 22.0,  # Total Investment
        7: 18.0   # Schedule Tier
    }
    for c_idx, w in col_w.items():
        ws1.column_dimensions[get_column_letter(c_idx)].width = w

    # Print Setup Sheet 1
    ws1.page_setup.orientation = ws1.ORIENTATION_PORTRAIT
    ws1.page_setup.paperSize = ws1.PAPERSIZE_LETTER
    ws1.sheet_properties.pageSetUpPr.fitToPage = True
    ws1.page_setup.fitToWidth = 1
    ws1.page_setup.fitToHeight = 1
    ws1.page_margins.left = 0.35
    ws1.page_margins.right = 0.35
    ws1.page_margins.top = 0.35
    ws1.page_margins.bottom = 0.35

    wb.save(target_path)
    print(f"SUCCESS: Saved Master CSI 01 74 23 Excel Bid Form to {target_path}")

if __name__ == "__main__":
    create_gc_bid_excel()
